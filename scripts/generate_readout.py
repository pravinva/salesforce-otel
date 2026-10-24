#!/usr/bin/env python3
"""Suncorp architecture readout (PDF + DOCX) in the Databricks Field Engineering addendum format."""

from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import Image, KeepTogether, Paragraph, Preformatted, SimpleDocTemplate, Table, TableStyle

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
BRICKSMITH = Path("/Users/pravin.varma/Documents/Demo/bricksmith")
PDF_OUT = DOCS / "Suncorp_Agentforce_OTel_Databricks_Readout.pdf"
DOCX_OUT = DOCS / "Suncorp_Agentforce_OTel_Databricks_Readout.docx"
FIG_ARCH = DOCS / "diagrams" / "suncorp_agentforce_otel_zerobus.png"
DB_LOGO = BRICKSMITH / "logos/default/databricks-full.png"

HEADER_TEXT = "Suncorp Agentforce Observability  |  Architecture Readout"
AUTHOR = "Pravin Varma"

NAVY = colors.HexColor("#1B3139")
LAVA = colors.HexColor("#FF3621")
BODY = colors.HexColor("#374151")
MUTED = colors.HexColor("#64748B")
LINE = colors.HexColor("#D9E0E7")
PALE = colors.HexColor("#F5F7F9")
WHITE = colors.white
PAGE_W, PAGE_H = A4

NAVY_RGB = RGBColor(0x1B, 0x31, 0x39)
LAVA_RGB = RGBColor(0xFF, 0x36, 0x21)
BODY_RGB = RGBColor(0x37, 0x41, 0x51)
MUTED_RGB = RGBColor(0x64, 0x74, 0x8B)
WHITE_RGB = RGBColor(0xFF, 0xFF, 0xFF)
HEADER_FILL = "1B3139"
ROW_ALT = "F5F7F9"
LINE_HEX = "D9E0E7"

KICKER = "ARCHITECTURE READOUT"
TITLE = "Agentforce session traces to Databricks Zerobus"
SUBTITLE = (
    "How Salesforce Agentforce session traces reach Unity Catalog Delta tables: session export GET, "
    "Protobuf POST to Zerobus Ingest OTLP, and where AWS Private Link attaches for Suncorp."
)

META_ROWS = [
    ["Customer", "Suncorp"],
    ["Scope", "Land Agentforce session traces in Unity Catalog Delta via Zerobus Ingest OTLP"],
    ["Lab environment", "Salesforce OrgFarm orgfarm918; Databricks FEVM fevm-ot-pdm-semi (us-east-1); AWS account 332745928618, EKS agentforce-otel (ap-southeast-2)"],
    ["Prepared by", AUTHOR],
    ["Date", "24 October 2026"],
    ["Status", "Draft v0.1 for Suncorp architecture review"],
]

CONTRACT_ROWS = [
    ["Job", "Session system of record", "Lakehouse ingest"],
    ["Call", "Integration GETs OTLP JSON", "Integration POSTs OTLP Protobuf"],
    ["Unit", "One session id, 72-hour window", "Batches of spans, logs, and metrics"],
    ["Auth", "Salesforce OAuth 2.0 External Client App", "Databricks service principal OAuth, Unity Catalog authorization"],
]

COMPONENT_ROWS = [
    [
        "Agentforce + Data 360 Session Tracing",
        "Salesforce",
        "Runs the agent and stores prompts, tool calls, evaluation spans, metrics, and logs per session.",
    ],
    [
        "Session export API",
        "Salesforce",
        "GET /services/data/v66.0/einstein/audit/otel/{session-id} returns OTLP JSON for one session.",
    ],
    [
        "Session poller",
        "Laptop in the lab; Suncorp VPC in production",
        "Discovers finished sessions, GETs JSON, converts to Protobuf, fills empty spanIds, POSTs to the Collector or to Zerobus.",
    ],
    [
        "OpenTelemetry Collector",
        "EKS Fargate agentforce-otel",
        "Receives Protobuf on 4317/4318, batches, sets source attributes, refreshes OAuth, POSTs to Zerobus.",
    ],
    [
        "Zerobus Ingest OTLP",
        "Databricks workspace region",
        "Accepts Protobuf with x-databricks-zerobus-table-name and writes Delta tables.",
    ],
    [
        "Unity Catalog Delta tables",
        "Databricks",
        "agentforce_otel_spans, _logs, _metrics with VARIANT attributes. SQL and MLflow evals read these tables.",
    ],
]

PHASE_ROWS = [
    [
        "A1",
        "Live Agentforce session 01a16cb0-5079-71ce-a37e-9aee0cb3c332: 8 spans, public Zerobus HTTP 200, rows in Delta.",
        "Done, 24 October 2026",
    ],
    [
        "A2",
        "OpenTelemetry Collector in an AWS VPC (lab host: EKS Fargate), POSTing to public Zerobus.",
        "Done, 24 October 2026",
    ],
    [
        "B",
        "Same Collector with Private Link: private DNS, VPC endpoint, POST and SQL from the VPC.",
        "Suncorp network build",
    ],
]

OBSERVE_ROWS = [
    [
        "agentforce_otel_spans",
        "Every successful trace POST",
        "New time; names such as GeneralCRM_*, AiCopilot__ReactInitialPrompt, CLOSED_USER_REQUEST, TaskResolution, InstructionAdherence; populated trace_id and span_id; status.code; attributes VARIANT with prompt and tool payload.",
    ],
    [
        "agentforce_otel_logs",
        "When the session export includes log records",
        "Rows sharing the same trace_id; body and severity_text.",
    ],
    [
        "agentforce_otel_metrics",
        "When the session export includes metrics",
        "Rows in the same window; metric name and points in VARIANT.",
    ],
]

SQL_RECENT = """SELECT time, service_name, name, status.code, trace_id, span_id
FROM ot_pdm_semi_catalog.agentforce_otel.agentforce_otel_spans
ORDER BY time DESC
LIMIT 20;"""

SQL_NAMES = """SELECT name, COUNT(*) AS spans
FROM ot_pdm_semi_catalog.agentforce_otel.agentforce_otel_spans
WHERE time > current_timestamp() - INTERVAL 1 DAYS
GROUP BY name
ORDER BY spans DESC;"""

LAB_PROD_ROWS = [
    ["Salesforce", "OrgFarm DE orgfarm918", "Suncorp Agentforce + Data Cloud"],
    ["Databricks", "Public FEVM fevm-ot-pdm-semi, us-east-1", "Private Link workspace in the same region as the VPC"],
    ["VPC", "Demo VPC with EKS agentforce-otel, ap-southeast-2", "Suncorp VPC; Collector uses the Private Link hostname"],
    ["Collector", "Running on Fargate", "Same pattern with private DNS and Private Link"],
    ["Poller", "Laptop with kubectl port-forward", "Same VPC as the Collector"],
    ["Catalog", "ot_pdm_semi_catalog.agentforce_otel", "Suncorp Unity Catalog names, same DDL"],
    ["Auth", "SP agentforce-zerobus-otel", "Suncorp SP with table-level SELECT and MODIFY"],
]

SECTIONS = {
    "purpose": (
        "This readout describes how Suncorp lands Salesforce Agentforce session traces in Databricks. "
        "It sets out the Salesforce session export API, the Zerobus Ingest OTLP contract, the telemetry "
        "bridge that connects them, and the point where AWS Private Link attaches. The test plan for the "
        "Private Link phase sits in the companion document suncorp-test-setup.md."
    ),
    "def_export": (
        "<b>Session export API</b> is Salesforce's Export Agentforce Session Tracing Data endpoint. A client "
        "calls GET with one session id and receives OTLP-shaped JSON for that session."
    ),
    "def_zerobus": (
        "<b>Zerobus Ingest OTLP</b> is the Databricks managed gateway that accepts OTLP Protobuf over HTTPS "
        "and writes Unity Catalog Delta tables."
    ),
    "def_bridge": (
        "<b>Telemetry bridge</b> is the poller plus a Zerobus client. The poller is the Salesforce client. "
        "The Zerobus client is the poller itself (Phase A1) or an OpenTelemetry Collector in the VPC (Phase A2). "
        "This lab hosted the Collector on EKS Fargate."
    ),
    "background_1": (
        "Suncorp runs customer-facing work on Salesforce Agentforce: LLM reasoning, tool calls, and CRM actions. "
        "Platform and risk teams query those traces in Databricks, run evaluations, and retain them with the rest "
        "of the lakehouse. Suncorp workspaces use AWS front-end Private Link. The ingest path is Zerobus Ingest "
        "OTLP into Delta. Delay follows Data 360 catch-up and the poll interval, typically tens of seconds to a "
        "few minutes after the session is readable in Salesforce."
    ),
    "background_2": (
        "Both products publish the APIs this design uses. This lab wires them together on Databricks and AWS "
        "and identifies the network work Suncorp owns for the Private Link phase."
    ),
    "salesforce_1": (
        "With Agentforce Session Tracing, Einstein Audit, and Agent Platform Tracing enabled, Data 360 stores "
        "each session: prompts, tool calls, evaluation spans, metrics, and logs. Salesforce's Export Agentforce "
        "Session Tracing Data API returns that payload over REST."
    ),
    "salesforce_2": (
        "A client authenticates with OAuth 2.0 through an External Client App, supplies one session id, and "
        "receives JSON that follows the OTLP schema (ResourceSpans, plus metrics and logs). Sessions must have "
        "started within the prior 72 hours. Throughput follows Connect API rate limits. Salesforce's developer "
        "guidance is to POST this JSON to the observability backend you operate."
    ),
    "salesforce_3": (
        "Data 360 DMOs such as ssot__TelemetryTraceSpan__dlm offer a second read path through SOQL. "
        "OTLP is the OpenTelemetry Protocol for traces, metrics, and logs over HTTP or gRPC. Zerobus Ingest OTLP "
        "accepts that protocol as Protobuf on POST. The Salesforce endpoint returns OTLP-shaped JSON on GET."
    ),
    "zerobus_1": (
        "Zerobus Ingest OTLP is a managed gateway. A client POSTs OTLP to "
        "{workspaceId}.zerobus.{region}.cloud.databricks.com with Unity Catalog authorization and the header "
        "x-databricks-zerobus-table-name. Zerobus writes Delta tables and stores attributes as VARIANT."
    ),
    "zerobus_2": (
        "The gateway takes Protobuf (application/x-protobuf). The client initiates every POST. The OTLP "
        "interface is a documented front-end Private Link service on TCP 443; MQTT and Kafka-compatible Zerobus "
        "APIs follow their own networking model. Databricks OAuth tokens last about 60 minutes, so the client "
        "refreshes them."
    ),
    "bridge_1": (
        "Salesforce holds the session and serves it on GET. Zerobus accepts a Protobuf POST and writes Delta. "
        "The telemetry bridge is the client of both APIs. Table 1 sets the two contracts side by side."
    ),
    "bridge_2": (
        "The bridge discovers finished sessions through ssot__AiAgentSession__dlm, reads each session with the "
        "export GET, converts OTLP JSON to OTLP Protobuf, and POSTs to Zerobus with a refreshed service principal "
        "token. Where Salesforce sends an empty spanId on evaluation spans, the converter writes a deterministic "
        "8-byte id so the batch is valid OTLP."
    ),
    "bridge_3": (
        "The poller is the Salesforce client. It can POST Protobuf straight to Zerobus (Phase A1). The Collector "
        "is the standing Databricks client for token refresh, batching, and Private Link egress from a VPC "
        "(Phase A2). Production can host that client on EKS, ECS, a small VM, or a scheduled function."
    ),
    "arch_1": (
        "Figure 1 reads left to right. Agentforce writes sessions into Data 360. The poller GETs the session JSON "
        "(PULL GET OTel JSON). The poller POSTs Protobuf to the Collector (PUSH Protobuf), and the Collector POSTs "
        "to Zerobus (PUSH OTLP Protobuf to Zerobus). Zerobus returns HTTP 200 once the write is durable. Phase A1 "
        "used the same Protobuf POST from the laptop with the Collector omitted."
    ),
    "arch_2": (
        "Today the Collector POSTs over public HTTPS to FEVM Zerobus in us-east-1. Phase B keeps the same "
        "Collector and Protobuf and points the hostname at a VPC endpoint: front-end or service-direct Private "
        "Link on TCP 443 with private DNS. The Salesforce GET uses Salesforce's network path. If that hop must "
        "also stay private, Suncorp adds Salesforce Private Connect or an approved egress allowlist as a "
        "separate control."
    ),
    "lab_1": (
        "A1 posted Protobuf from the laptop to public Zerobus. A2 runs the Collector in a VPC on the same public "
        "hostname. FEVM fevm-ot-pdm-semi is a public workspace. Suncorp copies the Collector-in-VPC pattern onto "
        "Private Link. Manifests are in otel-collector/k8s/."
    ),
    "lab_2": (
        "Phase B checks cover private DNS resolution, VPC endpoint connectivity, an OTLP POST returning HTTP 200, "
        "and SQL rows for the test session. The detailed cases sit in suncorp-test-setup.md."
    ),
    "delivery": (
        "Suncorp places a Zerobus client (Collector or equivalent) in its VPC, swaps workspace id, catalog, and "
        "schema, and keeps Protobuf with x-databricks-zerobus-table-name. Suncorp networking provisions the VPC "
        "endpoint and private DNS for Zerobus OTLP. The poller runs in the same VPC and authenticates to Salesforce "
        "through an External Client App. VARIANT columns keep prompt and tool payloads available for MLflow "
        "evaluation. Delta Sharing can return curated aggregates to Data Cloud once traces land in Delta."
    ),
    "run_1": (
        "Catalog ot_pdm_semi_catalog.agentforce_otel. Use the Suncorp Pilot SQL warehouse. DDL is sql/01_create_otel_tables.sql. "
        "Repeatable queries are sql/02_demo_queries.sql."
    ),
    "run_2": (
        "Produce a session: enable Session Tracing and Einstein Audit (dataspace default), then chat with Agentforce (Default) "
        "or run sf agent test run --api-name OTel_Demo_Test --target-org orgfarm918 --wait 10 --verbose. Lab utterance: "
        "Hello, what can you do? Expected topic: GeneralCRM. Copy generatedData.sessionId."
    ),
    "run_3": (
        "Wait for Data 360 catch-up (often several minutes). Run DRY_RUN=true python3 poller/poll_agentforce.py. "
        "Observe a session id, span count greater than zero, and complete status. Empty spans means Data 360 is still catching up."
    ),
    "run_4": (
        "POST Protobuf. Phase A1: set OTEL_COLLECTOR_TRACES_URL to "
        "https://{WORKSPACE_ID}.zerobus.{REGION}.cloud.databricks.com:443/v1/traces and run poller/poll_agentforce.py with "
        "SALESFORCE_SESSION_ID. Phase A2: kubectl port-forward svc/otel-collector 4318:4318 and POST to "
        "http://127.0.0.1:4318/v1/traces. Observe Exported N span(s) and HTTP 200. Synthetic check: "
        "python3 scripts/send_direct_zerobus.py."
    ),
    "run_5": (
        "A short Agentforce chat often fills spans first. Empty logs or metrics with populated spans is a valid session. "
        "Lab session 01a16cb0-5079-71ce-a37e-9aee0cb3c332 produced 8 spans and Zerobus HTTP 200. Match a new run by time "
        "descending and by name."
    ),
}

def pdf_styles():
    sample = getSampleStyleSheet()
    return {
        "kicker": ParagraphStyle("Kicker", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=9, leading=12, textColor=LAVA, spaceAfter=4),
        "title": ParagraphStyle("Title", parent=sample["Title"], fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=NAVY, alignment=TA_LEFT, spaceAfter=6),
        "subtitle": ParagraphStyle("Subtitle", parent=sample["BodyText"], fontName="Helvetica", fontSize=11, leading=15, textColor=MUTED, spaceAfter=10),
        "h1": ParagraphStyle("H1", parent=sample["Heading1"], fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=NAVY, spaceBefore=8, spaceAfter=6),
        "h2": ParagraphStyle("H2", parent=sample["Heading2"], fontName="Helvetica-Bold", fontSize=11.5, leading=15, textColor=NAVY, spaceBefore=8, spaceAfter=4),
        "body": ParagraphStyle("Body", parent=sample["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=BODY, alignment=TA_JUSTIFY, spaceAfter=7),
        "body_left": ParagraphStyle("BodyLeft", parent=sample["BodyText"], fontName="Helvetica", fontSize=9.5, leading=13, textColor=BODY, alignment=TA_LEFT, spaceAfter=7),
        "caption": ParagraphStyle("Caption", parent=sample["BodyText"], fontName="Helvetica-Oblique", fontSize=8, leading=11, textColor=MUTED, spaceBefore=3, spaceAfter=8),
        "table": ParagraphStyle("Table", parent=sample["BodyText"], fontName="Helvetica", fontSize=8, leading=10.5, textColor=BODY),
        "table_head": ParagraphStyle("TableHead", parent=sample["BodyText"], fontName="Helvetica-Bold", fontSize=8, leading=10.5, textColor=WHITE),
        "code": ParagraphStyle("Code", parent=sample["Code"], fontName="Courier", fontSize=7.6, leading=10, textColor=NAVY, backColor=PALE, borderPadding=6, spaceBefore=4, spaceAfter=10),
    }


def p(text, style):
    return Paragraph(text, style)


def header_footer(canvas, doc):
    canvas.saveState()
    canvas.setStrokeColor(LAVA)
    canvas.setLineWidth(1.4)
    canvas.line(18 * mm, PAGE_H - 16 * mm, PAGE_W - 18 * mm, PAGE_H - 16 * mm)
    if DB_LOGO.exists():
        canvas.drawImage(str(DB_LOGO), 18 * mm, PAGE_H - 12.4 * mm, width=28 * mm, height=6 * mm, preserveAspectRatio=True, mask="auto")
    canvas.setFillColor(NAVY)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawRightString(PAGE_W - 18 * mm, PAGE_H - 10.4 * mm, HEADER_TEXT)
    canvas.setFillColor(LAVA)
    canvas.rect(0, 0, PAGE_W, 1.6 * mm, stroke=0, fill=1)
    canvas.setFillColor(MUTED)
    canvas.setFont("Helvetica", 7.5)
    canvas.drawString(18 * mm, 6 * mm, AUTHOR)
    canvas.drawRightString(PAGE_W - 18 * mm, 6 * mm, f"Page {doc.page}")
    canvas.restoreState()


def data_table(headers, rows, widths, s, font_size=8.0):
    head = s["table_head"]
    body = ParagraphStyle(f"T{font_size}", parent=s["table"], fontSize=font_size, leading=font_size + 2.2)
    data = [[p(h, head) for h in headers]]
    data.extend([[p(str(c), body) for c in row] for row in rows])
    table = Table(data, colWidths=widths, repeatRows=1)
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), NAVY),
                ("GRID", (0, 0), (-1, -1), 0.4, LINE),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [WHITE, PALE]),
                ("VALIGN", (0, 0), (-1, -1), "TOP"),
                ("LEFTPADDING", (0, 0), (-1, -1), 5),
                ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                ("TOPPADDING", (0, 0), (-1, -1), 4.5),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 4.5),
            ]
        )
    )
    return table


def add_fig(story, path, caption, s, width=174 * mm):
    if not path.exists():
        return
    w, h = PILImage.open(path).size
    img = Image(str(path), width=width, height=width * h / w)
    img.hAlign = "CENTER"
    story.append(KeepTogether([img, p(caption, s["caption"])]))


def build_pdf():
    s = pdf_styles()
    story = [
        p(KICKER, s["kicker"]),
        p(TITLE, s["title"]),
        p(SUBTITLE, s["subtitle"]),
        data_table(["Field", "Value"], META_ROWS, [48 * mm, 126 * mm], s, 8.5),
        p("1. Purpose", s["h1"]),
        p(SECTIONS["purpose"], s["body"]),
        p("1.1 Key definitions", s["h2"]),
        p(SECTIONS["def_export"], s["body"]),
        p(SECTIONS["def_zerobus"], s["body"]),
        p(SECTIONS["def_bridge"], s["body"]),
        p("2. Background", s["h1"]),
        p(SECTIONS["background_1"], s["body"]),
        p(SECTIONS["background_2"], s["body"]),
        p("3. Salesforce session tracing", s["h1"]),
        p(SECTIONS["salesforce_1"], s["body"]),
        Preformatted("GET /services/data/v66.0/einstein/audit/otel/{session-id}", s["code"]),
        p(SECTIONS["salesforce_2"], s["body"]),
        p(SECTIONS["salesforce_3"], s["body"]),
        p("4. Databricks Zerobus Ingest", s["h1"]),
        p(SECTIONS["zerobus_1"], s["body_left"]),
        p(SECTIONS["zerobus_2"], s["body"]),
        p("5. The telemetry bridge", s["h1"]),
        p(SECTIONS["bridge_1"], s["body"]),
        data_table(["", "Salesforce session export API", "Databricks Zerobus OTLP"], CONTRACT_ROWS, [24 * mm, 75 * mm, 75 * mm], s, 7.8),
        p("Table 1. The two API contracts the bridge connects.", s["caption"]),
        p(SECTIONS["bridge_2"], s["body"]),
        p(SECTIONS["bridge_3"], s["body"]),
        p("6. Architecture", s["h1"]),
        p(SECTIONS["arch_1"], s["body"]),
        p(SECTIONS["arch_2"], s["body"]),
    ]
    add_fig(story, FIG_ARCH, "Figure 1. Agentforce session traces to Databricks: session export GET, telemetry bridge on EKS in an AWS VPC, Zerobus OTLP push, and the Phase B Private Link hop.", s, width=174 * mm)
    story += [
        data_table(["Component", "Where it runs", "Role"], COMPONENT_ROWS, [44 * mm, 38 * mm, 92 * mm], s, 7.6),
        p("Table 2. Components on the data path.", s["caption"]),
        KeepTogether([
            p("7. What the lab has demonstrated", s["h1"]),
            data_table(["Phase", "Work", "Status"], PHASE_ROWS, [16 * mm, 118 * mm, 40 * mm], s, 7.8),
            p("Table 3. Lab phases and status.", s["caption"]),
        ]),
        p(SECTIONS["lab_1"], s["body"]),
        KeepTogether([
            p("8. Lab and Suncorp production", s["h1"]),
            data_table(["", "This lab", "Suncorp production (Phase B)"], LAB_PROD_ROWS, [24 * mm, 70 * mm, 80 * mm], s, 7.8),
            p("Table 4. Lab environment against the Suncorp target.", s["caption"]),
        ]),
        p(SECTIONS["lab_2"], s["body"]),
        p("9. Produce traces and inspect tables", s["h1"]),
        p(SECTIONS["run_1"], s["body"]),
        p("9.1 Produce an Agentforce session", s["h2"]),
        p(SECTIONS["run_2"], s["body"]),
        p(SECTIONS["run_3"], s["body"]),
        p("9.2 POST to Zerobus", s["h2"]),
        p(SECTIONS["run_4"], s["body"]),
        p("9.3 Tables and what to look for", s["h2"]),
        data_table(["Table", "When it grows", "What to look for"], OBSERVE_ROWS, [38 * mm, 42 * mm, 94 * mm], s, 7.6),
        p("Table 5. Unity Catalog tables after a successful POST.", s["caption"]),
        p(SECTIONS["run_5"], s["body"]),
        Preformatted(SQL_RECENT, s["code"]),
        Preformatted(SQL_NAMES, s["code"]),
        p("10. Delivery implication", s["h1"]),
        p(SECTIONS["delivery"], s["body"]),
    ]

    doc = SimpleDocTemplate(str(PDF_OUT), pagesize=A4, leftMargin=18 * mm, rightMargin=18 * mm, topMargin=20 * mm, bottomMargin=14 * mm,
                            title=TITLE, author=AUTHOR)
    doc.build(story, onFirstPage=header_footer, onLaterPages=header_footer)
    print(f"Created {PDF_OUT}")


def strip_tags(text):
    return text.replace("<b>", "").replace("</b>", "")


def set_run_font(run, name="Calibri", size=11, bold=False, color=BODY_RGB):
    run.font.name = name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.font.color.rgb = color


def shade_cell(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), hex_color)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_border(cell, color=LINE_HEX):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tcPr.append(borders)


def add_heading(doc, text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        run.font.color.rgb = NAVY_RGB
        run.font.name = "Calibri"
        run.font.size = Pt(16 if level == 1 else 13)


def add_body(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(8)
    para.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    para.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if text.startswith("<b>"):
        term, rest = text[3:].split("</b>", 1)
        set_run_font(para.add_run(term), size=11, bold=True, color=BODY_RGB)
        set_run_font(para.add_run(rest), size=11, color=BODY_RGB)
    else:
        set_run_font(para.add_run(text), size=11, color=BODY_RGB)


def add_code(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(8)
    set_run_font(para.add_run(text), name="Courier New", size=8.5, color=NAVY_RGB)


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = ""
        set_run_font(cell.paragraphs[0].add_run(h), size=10, bold=True, color=WHITE_RGB)
        shade_cell(cell, HEADER_FILL)
        set_cell_border(cell)
    for r_idx, row in enumerate(rows):
        fill = ROW_ALT if r_idx % 2 else "FFFFFF"
        for c_idx, value in enumerate(row):
            cell = table.rows[r_idx + 1].cells[c_idx]
            cell.text = ""
            set_run_font(cell.paragraphs[0].add_run(str(value)), size=9.5, color=BODY_RGB)
            shade_cell(cell, fill)
            set_cell_border(cell)
    doc.add_paragraph()


def add_caption(doc, text):
    para = doc.add_paragraph()
    para.paragraph_format.space_after = Pt(12)
    run = para.add_run(text)
    set_run_font(run, size=9, color=MUTED_RGB)
    run.italic = True


def add_fig_docx(doc, path, caption):
    if not path.exists():
        return
    pic = doc.add_paragraph()
    pic.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic.add_run().add_picture(str(path), width=Inches(6.5))
    add_caption(doc, caption)


def add_border(paragraph, edge):
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    el = OxmlElement(f"w:{edge}")
    el.set(qn("w:val"), "single")
    el.set(qn("w:sz"), "12")
    el.set(qn("w:space"), "4")
    el.set(qn("w:color"), "FF3621")
    pBdr.append(el)
    pPr.append(pBdr)


def set_header_footer_docx(doc):
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(2.2)
    section.right_margin = Cm(2.2)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)

    hp = section.header.paragraphs[0]
    hp.clear()
    if DB_LOGO.exists():
        hp.add_run().add_picture(str(DB_LOGO), width=Inches(1.15))
    set_run_font(hp.add_run(f"    {HEADER_TEXT}"), size=9, color=NAVY_RGB)
    add_border(hp, "bottom")

    fp = section.footer.paragraphs[0]
    fp.clear()
    set_run_font(fp.add_run(AUTHOR), size=8, color=MUTED_RGB)
    fp.add_run("\t")
    runs = [fp.add_run(), fp.add_run(), fp.add_run()]
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    runs[0]._r.append(begin)
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    runs[1]._r.append(instr)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    runs[2]._r.append(end)
    for r in runs:
        set_run_font(r, size=8, color=MUTED_RGB)
    fp.paragraph_format.tab_stops.add_tab_stop(Inches(6.5), WD_TAB_ALIGNMENT.RIGHT)
    add_border(fp, "top")


def build_docx():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    style.font.color.rgb = BODY_RGB
    set_header_footer_docx(doc)

    kicker = doc.add_paragraph()
    set_run_font(kicker.add_run(KICKER), size=11, bold=True, color=LAVA_RGB)
    kicker.paragraph_format.space_after = Pt(4)
    title = doc.add_paragraph()
    set_run_font(title.add_run(TITLE), size=24, bold=True, color=NAVY_RGB)
    title.paragraph_format.space_after = Pt(8)
    add_body(doc, SUBTITLE)
    add_table(doc, ["Field", "Value"], META_ROWS)

    add_heading(doc, "1. Purpose", 1)
    add_body(doc, SECTIONS["purpose"])
    add_heading(doc, "1.1 Key definitions", 2)
    for key in ("def_export", "def_zerobus", "def_bridge"):
        add_body(doc, SECTIONS[key])

    add_heading(doc, "2. Background", 1)
    add_body(doc, SECTIONS["background_1"])
    add_body(doc, SECTIONS["background_2"])

    add_heading(doc, "3. Salesforce session tracing", 1)
    add_body(doc, SECTIONS["salesforce_1"])
    add_code(doc, "GET /services/data/v66.0/einstein/audit/otel/{session-id}")
    add_body(doc, SECTIONS["salesforce_2"])
    add_body(doc, SECTIONS["salesforce_3"])

    add_heading(doc, "4. Databricks Zerobus Ingest", 1)
    add_body(doc, SECTIONS["zerobus_1"])
    add_body(doc, SECTIONS["zerobus_2"])

    add_heading(doc, "5. The telemetry bridge", 1)
    add_body(doc, SECTIONS["bridge_1"])
    add_table(doc, ["", "Salesforce session export API", "Databricks Zerobus OTLP"], CONTRACT_ROWS)
    add_caption(doc, "Table 1. The two API contracts the bridge connects.")
    add_body(doc, SECTIONS["bridge_2"])
    add_body(doc, SECTIONS["bridge_3"])

    add_heading(doc, "6. Architecture", 1)
    add_body(doc, SECTIONS["arch_1"])
    add_body(doc, SECTIONS["arch_2"])
    add_fig_docx(doc, FIG_ARCH, "Figure 1. Agentforce session traces to Databricks: session export GET, telemetry bridge on EKS in an AWS VPC, Zerobus OTLP push, and the Phase B Private Link hop.")
    add_table(doc, ["Component", "Where it runs", "Role"], COMPONENT_ROWS)
    add_caption(doc, "Table 2. Components on the data path.")

    add_heading(doc, "7. What the lab has demonstrated", 1)
    add_table(doc, ["Phase", "Work", "Status"], PHASE_ROWS)
    add_caption(doc, "Table 3. Lab phases and status.")
    add_body(doc, SECTIONS["lab_1"])

    add_heading(doc, "8. Lab and Suncorp production", 1)
    add_table(doc, ["", "This lab", "Suncorp production (Phase B)"], LAB_PROD_ROWS)
    add_caption(doc, "Table 4. Lab environment against the Suncorp target.")
    add_body(doc, SECTIONS["lab_2"])

    add_heading(doc, "9. Produce traces and inspect tables", 1)
    add_body(doc, SECTIONS["run_1"])
    add_heading(doc, "9.1 Produce an Agentforce session", 2)
    add_body(doc, SECTIONS["run_2"])
    add_body(doc, SECTIONS["run_3"])
    add_heading(doc, "9.2 POST to Zerobus", 2)
    add_body(doc, SECTIONS["run_4"])
    add_heading(doc, "9.3 Tables and what to look for", 2)
    add_table(doc, ["Table", "When it grows", "What to look for"], OBSERVE_ROWS)
    add_caption(doc, "Table 5. Unity Catalog tables after a successful POST.")
    add_body(doc, SECTIONS["run_5"])
    add_code(doc, SQL_RECENT)
    add_code(doc, SQL_NAMES)

    add_heading(doc, "10. Delivery implication", 1)
    add_body(doc, SECTIONS["delivery"])

    doc.save(str(DOCX_OUT))
    print(f"Created {DOCX_OUT}")


if __name__ == "__main__":
    build_pdf()
    build_docx()
